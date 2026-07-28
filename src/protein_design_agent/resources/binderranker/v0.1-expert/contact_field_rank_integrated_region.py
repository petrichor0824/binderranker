#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
contact_field_rank_integrated_region.py

Integrated RFdiffusion binder ranking pipeline:
    PDB backbones -> V3 contact-field descriptors -> V4 morphology-adaptive ranking
    + weak design-region / hotspot-aware scoring.

================================================================================
需要手动修改/传参的地方（最重要）
================================================================================
1) 原始骨架路径：
   - 批量 PDB 文件夹：--input_dir ./your_pdb_folder
   - 单个 PDB 文件：  --input ./one_design.pdb
   - 已经有 V3 CSV 时：--input_csv ./contact_field_v3.csv

2) binder 链：
   - 默认 binder_chain=A，即 A 链是 binder，所有非 A 链都作为 target。
   - 如果你的 binder 不是 A 链，改：--binder_chain B

3) design region / desired region：
   - 如果你希望 binder 主要结合 target 某一区域，传：
       --desired_regions "B:45-60,B:80-95"
   - 格式支持：
       B:45-60   表示 PDB 中 B 链 45 到 60 号残基
       B:80      表示 B 链 80 号残基
       10-30     表示 target 列表中的第 10 到 30 个残基（不推荐多链时用）
   - 如果 target 已经是裁剪后的目标片段，可以不传 desired_regions。

4) undesired region：
   - 如果有不希望 binder 贴到的区域，传：
       --undesired_regions "B:150-180"
   - 没有就留空。

5) hotspot：
   - 如果 RFdiffusion 设计时用过 hotspot，建议传入同一批 hotspot：
       --hotspot_regions "B:48,B:52,B:56"
   - 脚本会同时算精确 hotspot 和 hotspot 邻域。
   - hotspot 邻域半径默认 8 Å，可改：
       --hotspot_expand_radius 8.0
   - 注意：hotspot 在最终分数中是弱加分/辅助项，不会压过整体几何质量。

6) 输出前缀：
   - --output_prefix ./results/qbqc_integrated
   - 会输出：
       ./results/qbqc_integrated_metrics.csv
       ./results/qbqc_integrated_scored.csv
       ./results/qbqc_integrated_ranking.xlsx
       ./results/qbqc_integrated_report.txt

推荐运行示例：

python contact_field_rank_integrated_region.py \
  --input_dir ./pdbs \
  --output_prefix ./results/qbqc_integrated \
  --binder_chain A \
  --desired_regions "B:45-113" \
  --undesired_regions "B:160-220" \
  --hotspot_regions "B:52,B:56,B:60" \
  --hotspot_expand_radius 8.0 \
  --reference_names "QbQc_219_PDs_4978,QbQc_219_PDs_3722"

核心设计原则：
- 原 V4 逻辑不删除：保留 final_score_v4_original。
- 新增 region/hotspot 弱加分：score_region, score_hotspot, final_score_v4_region。
- 默认主分数 final_score_v4 = final_score_v4_region（如果没有 region/hotspot，则自动等于原始分数）。
- region/hotspot 不做强硬过滤，除非你显式打开 --region_filter soft。

Author: ChatGPT, integrated from the user's contact_field_v3.py and rank_candidates_v4_layered.py.
"""

from __future__ import annotations

import argparse
import csv
import math
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import numpy as np

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
except Exception:
    Workbook = None

EPS = 1e-8
EPS_RANK = 1e-12


# =============================================================================
# V3: PDB parsing and contact-field descriptor extraction
# =============================================================================

class Residue:
    def __init__(self, chain: str, resid: int, resname: str, icode: str = ""):
        self.chain = chain if chain else "_"
        self.resid = resid
        self.resname = resname
        self.icode = icode.strip()
        self.atoms: Dict[str, np.ndarray] = {}

    def has(self, atom: str) -> bool:
        return atom in self.atoms

    @property
    def ca(self) -> np.ndarray:
        return self.atoms["CA"]

    @property
    def label(self) -> str:
        return f"{self.chain}:{self.resid}{self.icode}"


def parse_pdb_backbone(pdb_path: Path, binder_chain: str = "A") -> Tuple[List[Residue], List[Residue]]:
    """Read backbone atoms. Binder = binder_chain; target = all non-binder chains."""
    residues: Dict[Tuple[str, int, str, str], Residue] = {}
    with open(pdb_path, "r", encoding="utf-8", errors="ignore") as f:
        for line in f:
            if not line.startswith("ATOM"):
                continue
            atom = line[12:16].strip()
            if atom not in {"N", "CA", "C", "O", "CB"}:
                continue
            resname = line[17:20].strip()
            chain = line[21].strip() or "_"
            icode = line[26].strip()
            try:
                resid = int(line[22:26])
                x = float(line[30:38])
                y = float(line[38:46])
                z = float(line[46:54])
            except ValueError:
                continue
            key = (chain, resid, icode, resname)
            if key not in residues:
                residues[key] = Residue(chain, resid, resname, icode)
            residues[key].atoms[atom] = np.array([x, y, z], dtype=float)

    binder, target = [], []
    for r in residues.values():
        if not r.has("CA"):
            continue
        if r.chain == binder_chain:
            binder.append(r)
        else:
            target.append(r)

    binder.sort(key=lambda r: (r.chain, r.resid, r.icode))
    target.sort(key=lambda r: (r.chain, r.resid, r.icode))

    if not binder:
        raise ValueError(f"No binder residues found for binder_chain={binder_chain}")
    if not target:
        raise ValueError("No target residues found. Target = all non-binder chains.")
    return binder, target


def sigmoid(x: np.ndarray) -> np.ndarray:
    x = np.clip(x, -60, 60)
    return 1.0 / (1.0 + np.exp(-x))


def pairwise_dist(a: np.ndarray, b: np.ndarray) -> np.ndarray:
    diff = a[:, None, :] - b[None, :, :]
    return np.linalg.norm(diff, axis=-1)


def pseudo_cb_coord(res: Residue) -> Optional[np.ndarray]:
    if res.has("CB"):
        return res.atoms["CB"]
    if not (res.has("N") and res.has("CA") and res.has("C")):
        return None
    n, ca, c = res.atoms["N"], res.atoms["CA"], res.atoms["C"]
    a = n - ca
    b = c - ca
    cvec = np.cross(a, b)
    cb = ca + (-0.58273431 * a + 0.56802827 * b - 0.54067466 * cvec)
    return cb


def pseudo_cb_direction(res: Residue) -> Tuple[Optional[np.ndarray], Optional[np.ndarray]]:
    cb = pseudo_cb_coord(res)
    if cb is None:
        return None, None
    u = cb - res.ca
    norm = np.linalg.norm(u)
    if norm < EPS:
        return cb, None
    return cb, u / norm


def residue_arrays(residues: List[Residue]) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    ca = np.array([r.ca for r in residues], dtype=float)
    cb_list, dir_list = [], []
    for r in residues:
        cb, u = pseudo_cb_direction(r)
        if cb is None or u is None:
            cb_list.append(r.ca.copy())
            dir_list.append(np.array([np.nan, np.nan, np.nan], dtype=float))
        else:
            cb_list.append(cb)
            dir_list.append(u)
    return ca, np.array(cb_list, dtype=float), np.array(dir_list, dtype=float)


def has_valid_dir(dirs: np.ndarray) -> np.ndarray:
    return ~np.isnan(dirs).any(axis=1)


def effective_number(x: np.ndarray) -> float:
    s = float(np.sum(x))
    if s <= EPS:
        return 0.0
    return (s * s) / (float(np.sum(x * x)) + EPS)


def entropy_norm(x: np.ndarray) -> float:
    s = float(np.sum(x))
    if s <= EPS:
        return 0.0
    p = x / s
    p = p[p > EPS]
    if len(p) <= 1:
        return 0.0
    return -float(np.sum(p * np.log(p))) / math.log(len(x))


def max_pairwise_distance(coords: np.ndarray) -> float:
    if len(coords) < 2:
        return 0.0
    return float(np.max(pairwise_dist(coords, coords)))


def pca_shape_features(coords: np.ndarray) -> Tuple[float, float, float]:
    if coords.shape[0] < 3:
        return 0.0, 0.0, 0.0
    x = coords - np.mean(coords, axis=0, keepdims=True)
    cov = x.T @ x / max(coords.shape[0] - 1, 1)
    vals, _ = np.linalg.eigh(cov)
    vals = np.sort(vals)[::-1]
    l1, l2, l3 = vals[0], vals[1], vals[2]
    if l1 <= EPS:
        return 0.0, 0.0, 0.0
    linearity = float((l1 - l2) / (l1 + EPS))
    planarity = float((l2 - l3) / (l1 + EPS))
    compactness = float(l3 / (l1 + EPS))
    return max(0.0, linearity), max(0.0, planarity), max(0.0, compactness)


def pca_plane_normal(coords: np.ndarray) -> Tuple[Optional[np.ndarray], float]:
    if coords.shape[0] < 3:
        return None, 0.0
    x = coords - np.mean(coords, axis=0, keepdims=True)
    cov = x.T @ x / max(coords.shape[0] - 1, 1)
    vals, vecs = np.linalg.eigh(cov)
    order = np.argsort(vals)[::-1]
    vals = vals[order]
    vecs = vecs[:, order]
    normal = vecs[:, -1]
    normal = normal / (np.linalg.norm(normal) + EPS)
    planarity_conf = 0.0 if vals[1] <= EPS else 1.0 - float(vals[-1] / (vals[1] + EPS))
    return normal, max(0.0, min(1.0, planarity_conf))


def parse_region_specs(spec: str) -> List[str]:
    if spec is None:
        return []
    spec = spec.strip()
    if not spec:
        return []
    return [x.strip() for x in spec.split(",") if x.strip()]


def region_mask_for_target(target: List[Residue], specs: List[str]) -> np.ndarray:
    """
    Build target-region mask.
    Supported formats:
      B:45-60, B:80      real PDB chain/residue numbering
      10-30, 15          target-list 1-based positions
    """
    mask = np.zeros(len(target), dtype=bool)
    for sp in specs:
        try:
            if ":" in sp:
                chain, rr = sp.split(":", 1)
                chain = chain.strip()
                if "-" in rr:
                    a, b = rr.split("-", 1)
                    start, end = int(a), int(b)
                else:
                    start = end = int(rr)
                if end < start:
                    start, end = end, start
                for i, r in enumerate(target):
                    if r.chain == chain and start <= r.resid <= end:
                        mask[i] = True
            else:
                if "-" in sp:
                    a, b = sp.split("-", 1)
                    start, end = int(a), int(b)
                else:
                    start = end = int(sp)
                if end < start:
                    start, end = end, start
                start_idx = max(1, start)
                end_idx = min(len(target), end)
                if start_idx <= end_idx:
                    mask[start_idx - 1:end_idx] = True
        except Exception:
            # Keep robust: an invalid region spec should not crash an entire batch.
            continue
    return mask


def expand_mask_by_radius(coords: np.ndarray, seed_mask: np.ndarray, radius: float) -> np.ndarray:
    """Expand a target region by CA distance radius."""
    seed_mask = np.array(seed_mask, dtype=bool)
    if len(coords) == 0 or not np.any(seed_mask) or radius <= 0:
        return seed_mask.copy()
    seed_coords = coords[seed_mask]
    d = pairwise_dist(coords, seed_coords)
    return np.min(d, axis=1) <= radius


def weighted_mean(values: np.ndarray, weights: np.ndarray) -> float:
    w = float(np.sum(weights))
    if w <= EPS:
        return 0.0
    return float(np.sum(values * weights) / (w + EPS))


def build_spatial_edges(coords: np.ndarray, cutoff: float, dirs: Optional[np.ndarray] = None,
                        face_aware: bool = False, face_floor: float = 0.35) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    n = len(coords)
    if n < 2:
        return np.array([], dtype=int), np.array([], dtype=int), np.array([], dtype=float)
    D = pairwise_dist(coords, coords)
    mask = np.triu((D < cutoff) & (D > EPS), k=1)
    edge_i, edge_j = np.where(mask)
    if len(edge_i) == 0:
        return edge_i, edge_j, np.array([], dtype=float)
    d = D[edge_i, edge_j]
    spatial_w = np.exp(-(d ** 2) / (2.0 * (cutoff / 2.0) ** 2))
    if face_aware and dirs is not None:
        valid = has_valid_dir(dirs)
        both_valid = valid[edge_i] & valid[edge_j]
        face_w = np.ones(len(edge_i), dtype=float)
        if np.any(both_valid):
            dots = np.sum(dirs[edge_i[both_valid]] * dirs[edge_j[both_valid]], axis=1)
            dots = np.clip(dots, -1.0, 1.0)
            face_w[both_valid] = face_floor + (1.0 - face_floor) * np.clip(dots, 0.0, 1.0)
        edge_w = spatial_w * face_w
    else:
        edge_w = spatial_w
    return edge_i, edge_j, edge_w


def graph_field_roughness(coords: np.ndarray, field: np.ndarray, cutoff: float,
                          dirs: Optional[np.ndarray] = None, face_aware: bool = False,
                          contact_threshold: float = 0.1) -> Dict[str, float]:
    edge_i, edge_j, edge_w = build_spatial_edges(coords=coords, cutoff=cutoff, dirs=dirs, face_aware=face_aware)
    if len(edge_i) == 0:
        return {"active_roughness": 0.0, "active_smoothness": 0.0, "active_edges": 0,
                "core_roughness": 0.0, "core_smoothness": 0.0, "core_edges": 0}
    fi, fj = field[edge_i], field[edge_j]
    active_mask = (fi > contact_threshold) | (fj > contact_threshold)
    core_mask = (fi > contact_threshold) & (fj > contact_threshold)

    def calc(mask: np.ndarray) -> Tuple[float, float, int]:
        if not np.any(mask):
            return 0.0, 0.0, 0
        diff = np.abs(fi[mask] - fj[mask])
        avg = 0.5 * (fi[mask] + fj[mask]) + EPS
        w = edge_w[mask]
        rough = float(np.sum(w * diff) / (np.sum(w * avg) + EPS))
        return rough, math.exp(-rough), int(np.sum(mask))

    active_rough, active_smooth, active_edges = calc(active_mask)
    core_rough, core_smooth, core_edges = calc(core_mask)
    return {"active_roughness": active_rough, "active_smoothness": active_smooth, "active_edges": active_edges,
            "core_roughness": core_rough, "core_smoothness": core_smooth, "core_edges": core_edges}


def compute_contact_weights(binder: List[Residue], target: List[Residue], shell_max: float, ideal_dist: float,
                            density_sigma: float, clash_cutoff: float, tangent_floor: float, front_dot: float,
                            back_dot: float, dot_sigma: float, back_penalty: float, cb_far_delta: float,
                            cb_delta_sigma: float, cb_near_cutoff: float, cb_support_cutoff: float,
                            cb_sigma: float) -> Dict[str, np.ndarray]:
    b_ca, b_cb, b_dirs = residue_arrays(binder)
    t_ca, t_cb, t_dirs = residue_arrays(target)
    D = pairwise_dist(b_ca, t_ca)
    cb_D = pairwise_dist(b_cb, t_ca)
    raw_mask = D < shell_max
    dist_w = np.exp(-((D - ideal_dist) ** 2) / (2.0 * density_sigma * density_sigma)) * raw_mask
    dist_w[D < clash_cutoff] *= 0.25

    v = t_ca[None, :, :] - b_ca[:, None, :]
    v_norm = np.linalg.norm(v, axis=-1, keepdims=True)
    v_unit = v / (v_norm + EPS)
    valid_b_dir = has_valid_dir(b_dirs)
    dot = np.sum(b_dirs[:, None, :] * v_unit, axis=-1)
    dot[~valid_b_dir, :] = 0.0
    cb_delta = cb_D - D

    front_score = sigmoid((dot - front_dot) / max(dot_sigma, EPS))
    orient_w = tangent_floor + (1.0 - tangent_floor) * front_score
    back_orientation_score = sigmoid((back_dot - dot) / max(dot_sigma, EPS))
    cb_far_score = sigmoid((cb_delta - cb_far_delta) / max(cb_delta_sigma, EPS))
    back_suspicion = back_orientation_score * cb_far_score * raw_mask
    orient_w = orient_w * (1.0 - (1.0 - back_penalty) * back_suspicion)
    orient_w[~valid_b_dir, :] = 1.0

    cb_support_score = sigmoid((cb_support_cutoff - cb_D) / max(cb_sigma, EPS))
    cb_support_w = 0.65 + 0.35 * cb_support_score
    W = dist_w * orient_w * cb_support_w

    return {
        "b_ca": b_ca, "t_ca": t_ca, "b_dirs": b_dirs, "t_dirs": t_dirs,
        "D": D, "cb_D": cb_D, "cb_delta": cb_delta, "raw_mask": raw_mask,
        "dist_w": dist_w, "W": W, "dot": dot,
        "front_mask": (dot >= front_dot) & raw_mask,
        "tangential_mask": (dot > back_dot) & (dot < front_dot) & raw_mask,
        "backfacing_mask": (dot <= back_dot) & raw_mask,
        "backfacing_cb_far_mask": (dot <= back_dot) & (cb_delta > cb_far_delta) & raw_mask,
        "cb_near_mask": (cb_D < cb_near_cutoff) & raw_mask,
        "cb_supported_mask": (cb_D < cb_support_cutoff) & raw_mask,
        "cb_closer_mask": (cb_delta < 0.0) & raw_mask,
    }


def weighted_ratio(base_weight: np.ndarray, mask: np.ndarray, denom: float) -> float:
    if denom <= EPS:
        return 0.0
    return float(np.sum(base_weight[mask]) / (denom + EPS))


def empty_microdomain_result() -> Dict[str, float]:
    return {"microdomain_count": 0, "micro_plane_alignment_weighted": 0.0,
            "micro_center_distance_weighted": 0.0, "micro_planarity_mean": 0.0,
            "micro_fit_score_weighted": 0.0}


def compute_microdomain_plane_metrics(W: np.ndarray, b_ca: np.ndarray, t_ca: np.ndarray, t_field: np.ndarray,
                                      contact_threshold: float, micro_radius: float, ideal_dist: float,
                                      micro_dist_sigma: float, max_microdomains: int) -> Dict[str, float]:
    if len(t_ca) < 3 or len(b_ca) < 3:
        return empty_microdomain_result()
    t_internal_D = pairwise_dist(t_ca, t_ca)
    contacted_indices = np.where(t_field > contact_threshold)[0]
    if len(contacted_indices) == 0:
        return empty_microdomain_result()
    if len(contacted_indices) > max_microdomains:
        order = np.argsort(t_field[contacted_indices])[::-1]
        contacted_indices = contacted_indices[order[:max_microdomains]]

    aligns, dists, plans, fits, weights = [], [], [], [], []
    for j in contacted_indices:
        target_micro_idx = np.where(t_internal_D[j] < micro_radius)[0]
        if len(target_micro_idx) < 3:
            continue
        local_b_weight = np.sum(W[:, target_micro_idx], axis=1)
        binder_micro_idx = np.where(local_b_weight > contact_threshold)[0]
        if len(binder_micro_idx) < 3:
            topk = min(8, len(b_ca))
            top_idx = np.argsort(local_b_weight)[::-1][:topk]
            binder_micro_idx = top_idx[local_b_weight[top_idx] > EPS]
        if len(binder_micro_idx) < 3:
            continue
        t_coords, b_coords = t_ca[target_micro_idx], b_ca[binder_micro_idx]
        nT, pT = pca_plane_normal(t_coords)
        nB, pB = pca_plane_normal(b_coords)
        if nT is None or nB is None:
            continue
        alignment = abs(float(np.dot(nT, nB)))
        center_dist = float(np.linalg.norm(np.mean(t_coords, axis=0) - np.mean(b_coords, axis=0)))
        planarity = pT * pB
        distance_score = math.exp(-((center_dist - ideal_dist) ** 2) / (2.0 * micro_dist_sigma * micro_dist_sigma))
        local_contact_weight = float(np.sum(local_b_weight[binder_micro_idx]))
        fit_score = alignment * distance_score * planarity
        aligns.append(alignment); dists.append(center_dist); plans.append(planarity); fits.append(fit_score)
        weights.append(local_contact_weight * max(planarity, 0.05))
    if len(weights) == 0 or np.sum(weights) <= EPS:
        return empty_microdomain_result()
    aligns, dists, plans, fits, weights = map(np.array, (aligns, dists, plans, fits, weights))
    wsum = float(np.sum(weights)) + EPS
    return {"microdomain_count": int(len(weights)),
            "micro_plane_alignment_weighted": float(np.sum(aligns * weights) / wsum),
            "micro_center_distance_weighted": float(np.sum(dists * weights) / wsum),
            "micro_planarity_mean": float(np.mean(plans)),
            "micro_fit_score_weighted": float(np.sum(fits * weights) / wsum)}


def empty_continuity_result() -> Dict[str, float]:
    return {"contact_map_edges": 0, "contact_map_continuity_mean_dist": 0.0,
            "contact_map_continuity_score": 0.0, "contact_map_distortion_mean": 0.0,
            "contact_map_jump_fraction": 0.0}


def compute_contact_map_continuity(W: np.ndarray, b_ca: np.ndarray, t_ca: np.ndarray, b_field: np.ndarray,
                                   b_dirs: np.ndarray, graph_cutoff: float, contact_threshold: float,
                                   jump_cutoff: float) -> Dict[str, float]:
    mapped_target = np.zeros_like(b_ca)
    for i in range(len(b_ca)):
        if b_field[i] > EPS:
            mapped_target[i] = np.sum(W[i, :, None] * t_ca, axis=0) / (b_field[i] + EPS)
    edge_i, edge_j, edge_w = build_spatial_edges(b_ca, graph_cutoff, dirs=b_dirs, face_aware=True)
    if len(edge_i) == 0:
        return empty_continuity_result()
    active = (b_field[edge_i] > contact_threshold) & (b_field[edge_j] > contact_threshold)
    if not np.any(active):
        return empty_continuity_result()
    ei, ej, ew = edge_i[active], edge_j[active], edge_w[active]
    mapped_dist = np.linalg.norm(mapped_target[ei] - mapped_target[ej], axis=1)
    binder_dist = np.linalg.norm(b_ca[ei] - b_ca[ej], axis=1)
    mean_mapped_dist = weighted_mean(mapped_dist, ew)
    mean_distortion = weighted_mean(np.abs(mapped_dist - binder_dist), ew)
    jump_fraction = float(np.sum(ew[mapped_dist > jump_cutoff]) / (np.sum(ew) + EPS))
    return {"contact_map_edges": int(len(ei)),
            "contact_map_continuity_mean_dist": mean_mapped_dist,
            "contact_map_continuity_score": math.exp(-mean_mapped_dist / 10.0),
            "contact_map_distortion_mean": mean_distortion,
            "contact_map_jump_fraction": jump_fraction}


def region_metrics(t_field: np.ndarray, mask: np.ndarray, total_weight: float,
                   contact_threshold: float, coords: Optional[np.ndarray] = None) -> Dict[str, float]:
    count = int(np.sum(mask))
    if count == 0:
        return {"count": 0, "effective_coverage": 0.0, "contact_count": 0,
                "weight_fraction": 0.0, "mean_field": 0.0,
                "neff_norm": 0.0, "entropy": 0.0, "span": 0.0, "span_norm_global": 0.0}
    region_field = t_field[mask]
    contact_count = int(np.sum(region_field > contact_threshold))
    region_coords = coords[mask] if coords is not None else np.zeros((0, 3), dtype=float)
    active_coords = region_coords[region_field > contact_threshold] if coords is not None else np.zeros((0, 3), dtype=float)
    global_span = max_pairwise_distance(coords) if coords is not None else 0.0
    return {
        "count": count,
        "effective_coverage": float(np.mean(region_field > contact_threshold)),
        "contact_count": contact_count,
        "weight_fraction": float(np.sum(region_field) / (total_weight + EPS)),
        "mean_field": float(np.mean(region_field)),
        "neff_norm": effective_number(region_field) / max(count, 1),
        "entropy": entropy_norm(region_field) if count > 1 else 0.0,
        "span": max_pairwise_distance(active_coords),
        "span_norm_global": max_pairwise_distance(active_coords) / (global_span + EPS) if global_span > EPS else 0.0,
    }


def compute_metrics(pdb_path: Path, args: argparse.Namespace,
                    desired_specs: List[str], undesired_specs: List[str], hotspot_specs: List[str]) -> Dict[str, Any]:
    binder, target = parse_pdb_backbone(pdb_path, binder_chain=args.binder_chain)
    n_b, n_t = len(binder), len(target)

    data = compute_contact_weights(
        binder=binder, target=target, shell_max=args.shell_max, ideal_dist=args.ideal_dist,
        density_sigma=args.density_sigma, clash_cutoff=args.clash_cutoff,
        tangent_floor=args.tangent_floor, front_dot=args.front_dot, back_dot=args.back_dot,
        dot_sigma=args.dot_sigma, back_penalty=args.back_penalty, cb_far_delta=args.cb_far_delta,
        cb_delta_sigma=args.cb_delta_sigma, cb_near_cutoff=args.cb_near_cutoff,
        cb_support_cutoff=args.cb_support_cutoff, cb_sigma=args.cb_sigma)

    b_ca, t_ca, b_dirs, t_dirs = data["b_ca"], data["t_ca"], data["b_dirs"], data["t_dirs"]
    D, cb_D, cb_delta = data["D"], data["cb_D"], data["cb_delta"]
    raw_mask, dist_w, W = data["raw_mask"], data["dist_w"], data["W"]

    raw_weight_sum = float(np.sum(dist_w))
    effective_weight_sum = float(np.sum(W))
    effective_to_raw_ratio = effective_weight_sum / (raw_weight_sum + EPS)

    b_field = np.sum(W, axis=1)
    t_field = np.sum(W, axis=0)

    contacted_b = b_ca[b_field > args.contact_threshold]
    contacted_t = t_ca[t_field > args.contact_threshold]

    binder_rough = graph_field_roughness(b_ca, b_field, args.graph_cutoff, dirs=b_dirs, face_aware=True,
                                         contact_threshold=args.contact_threshold)
    target_rough = graph_field_roughness(t_ca, t_field, args.graph_cutoff, dirs=t_dirs, face_aware=False,
                                         contact_threshold=args.contact_threshold)
    continuity = compute_contact_map_continuity(W, b_ca, t_ca, b_field, b_dirs, args.graph_cutoff,
                                                args.contact_threshold, args.jump_cutoff)

    target_linearity, target_planarity, target_compactness = pca_shape_features(t_ca)
    if len(contacted_t) >= 3:
        ct_linearity, ct_planarity, ct_compactness = pca_shape_features(contacted_t)
    else:
        ct_linearity, ct_planarity, ct_compactness = 0.0, 0.0, 0.0

    micro = compute_microdomain_plane_metrics(W, b_ca, t_ca, t_field, args.contact_threshold,
                                              args.micro_radius, args.ideal_dist,
                                              args.micro_dist_sigma, args.max_microdomains)

    desired_mask = region_mask_for_target(target, desired_specs)
    undesired_mask = region_mask_for_target(target, undesired_specs)
    hotspot_mask = region_mask_for_target(target, hotspot_specs)
    hotspot_neighborhood_mask = expand_mask_by_radius(t_ca, hotspot_mask, args.hotspot_expand_radius)

    total_t_weight = float(np.sum(t_field))
    desired = region_metrics(t_field, desired_mask, total_t_weight, args.contact_threshold, coords=t_ca)
    undesired = region_metrics(t_field, undesired_mask, total_t_weight, args.contact_threshold, coords=t_ca)
    hotspot = region_metrics(t_field, hotspot_mask, total_t_weight, args.contact_threshold, coords=t_ca)
    hotspot_nb = region_metrics(t_field, hotspot_neighborhood_mask, total_t_weight, args.contact_threshold, coords=t_ca)

    cb_delta_weighted_mean = float(np.sum(dist_w * cb_delta) / (raw_weight_sum + EPS)) if raw_weight_sum > EPS else 0.0
    cb_dist_weighted_mean = float(np.sum(dist_w * cb_D) / (raw_weight_sum + EPS)) if raw_weight_sum > EPS else 0.0

    row: Dict[str, Any] = {
        "pdb_name": pdb_path.stem, "binder_len": n_b, "target_len": n_t,
        "raw_mean_shell_count": float(np.mean(np.sum(raw_mask, axis=1))),
        "raw_mean_shell_6": float(np.mean(np.sum(D < 6.0, axis=1))),
        "raw_mean_shell_7": float(np.mean(np.sum(D < 7.0, axis=1))),
        "raw_mean_shell_8": float(np.mean(np.sum(D < 8.0, axis=1))),
        "raw_mean_shell_10": float(np.mean(np.sum(D < 10.0, axis=1))),
        "raw_mean_shell_12": float(np.mean(np.sum(D < 12.0, axis=1))),
        "raw_weight_sum": raw_weight_sum,
        "effective_weight_sum": effective_weight_sum,
        "effective_to_raw_ratio": effective_to_raw_ratio,
        "orientation_discount_ratio": max(0.0, min(1.0, 1.0 - effective_to_raw_ratio)),
        "front_facing_weight_ratio": weighted_ratio(dist_w, data["front_mask"], raw_weight_sum),
        "tangential_weight_ratio": weighted_ratio(dist_w, data["tangential_mask"], raw_weight_sum),
        "backfacing_weight_ratio": weighted_ratio(dist_w, data["backfacing_mask"], raw_weight_sum),
        "backfacing_cb_far_weight_ratio": weighted_ratio(dist_w, data["backfacing_cb_far_mask"], raw_weight_sum),
        "cb_near_weight_ratio": weighted_ratio(dist_w, data["cb_near_mask"], raw_weight_sum),
        "cb_supported_weight_ratio": weighted_ratio(dist_w, data["cb_supported_mask"], raw_weight_sum),
        "cb_closer_weight_ratio": weighted_ratio(dist_w, data["cb_closer_mask"], raw_weight_sum),
        "cb_delta_weighted_mean": cb_delta_weighted_mean,
        "cb_dist_weighted_mean": cb_dist_weighted_mean,
        "clash_pairs": int(np.sum(D < args.clash_cutoff)),
        "min_distance_global": float(np.min(D)),
        "binder_effective_hit_ratio": float(np.mean(b_field > args.contact_threshold)),
        "target_effective_coverage": float(np.mean(t_field > args.contact_threshold)),
        "binder_contact_neff_norm": effective_number(b_field) / n_b,
        "target_contact_neff_norm": effective_number(t_field) / n_t,
        "binder_contact_entropy": entropy_norm(b_field),
        "target_contact_entropy": entropy_norm(t_field),
        "binder_contact_span_norm": max_pairwise_distance(contacted_b) / (max_pairwise_distance(b_ca) + EPS),
        "target_contact_span": max_pairwise_distance(contacted_t),
        "target_contact_span_norm": max_pairwise_distance(contacted_t) / (max_pairwise_distance(t_ca) + EPS),
        "binder_field_active_roughness": binder_rough["active_roughness"],
        "binder_field_active_smoothness": binder_rough["active_smoothness"],
        "binder_field_active_edges": binder_rough["active_edges"],
        "binder_field_core_roughness": binder_rough["core_roughness"],
        "binder_field_core_smoothness": binder_rough["core_smoothness"],
        "binder_field_core_edges": binder_rough["core_edges"],
        "target_field_active_roughness": target_rough["active_roughness"],
        "target_field_active_smoothness": target_rough["active_smoothness"],
        "target_field_active_edges": target_rough["active_edges"],
        "target_field_core_roughness": target_rough["core_roughness"],
        "target_field_core_smoothness": target_rough["core_smoothness"],
        "target_field_core_edges": target_rough["core_edges"],
        **continuity,
        "target_linearity": target_linearity,
        "target_planarity": target_planarity,
        "target_compactness": target_compactness,
        "contacted_target_linearity": ct_linearity,
        "contacted_target_planarity": ct_planarity,
        "contacted_target_compactness": ct_compactness,
        **micro,
        # Original desired/undesired region metrics.
        "desired_region_count": desired["count"],
        "desired_contact_count": desired["contact_count"],
        "desired_effective_coverage": desired["effective_coverage"],
        "desired_weight_fraction": desired["weight_fraction"],
        "desired_mean_field": desired["mean_field"],
        "desired_contact_neff_norm": desired["neff_norm"],
        "desired_contact_entropy": desired["entropy"],
        "desired_contact_span": desired["span"],
        "desired_contact_span_norm_global": desired["span_norm_global"],
        "undesired_region_count": undesired["count"],
        "undesired_contact_count": undesired["contact_count"],
        "undesired_effective_coverage": undesired["effective_coverage"],
        "undesired_weight_fraction": undesired["weight_fraction"],
        "undesired_mean_field": undesired["mean_field"],
        "undesired_contact_neff_norm": undesired["neff_norm"],
        "undesired_contact_entropy": undesired["entropy"],
        "undesired_contact_span": undesired["span"],
        "undesired_contact_span_norm_global": undesired["span_norm_global"],
        # New hotspot metrics.
        "hotspot_region_count": hotspot["count"],
        "hotspot_contact_count": hotspot["contact_count"],
        "hotspot_effective_coverage": hotspot["effective_coverage"],
        "hotspot_weight_fraction": hotspot["weight_fraction"],
        "hotspot_mean_field": hotspot["mean_field"],
        "hotspot_contact_neff_norm": hotspot["neff_norm"],
        "hotspot_contact_entropy": hotspot["entropy"],
        "hotspot_neighborhood_radius": args.hotspot_expand_radius,
        "hotspot_neighborhood_count": hotspot_nb["count"],
        "hotspot_neighborhood_contact_count": hotspot_nb["contact_count"],
        "hotspot_neighborhood_effective_coverage": hotspot_nb["effective_coverage"],
        "hotspot_neighborhood_weight_fraction": hotspot_nb["weight_fraction"],
        "hotspot_neighborhood_mean_field": hotspot_nb["mean_field"],
        "hotspot_neighborhood_contact_neff_norm": hotspot_nb["neff_norm"],
        "hotspot_neighborhood_contact_entropy": hotspot_nb["entropy"],
        "hotspot_neighborhood_contact_span": hotspot_nb["span"],
        "hotspot_neighborhood_contact_span_norm_global": hotspot_nb["span_norm_global"],
    }
    return row


FIELDNAMES = [
    "pdb_name", "binder_len", "target_len",
    "raw_mean_shell_count", "raw_mean_shell_6", "raw_mean_shell_7", "raw_mean_shell_8", "raw_mean_shell_10", "raw_mean_shell_12",
    "raw_weight_sum", "effective_weight_sum", "effective_to_raw_ratio", "orientation_discount_ratio",
    "front_facing_weight_ratio", "tangential_weight_ratio", "backfacing_weight_ratio", "backfacing_cb_far_weight_ratio",
    "cb_near_weight_ratio", "cb_supported_weight_ratio", "cb_closer_weight_ratio", "cb_delta_weighted_mean", "cb_dist_weighted_mean",
    "clash_pairs", "min_distance_global",
    "binder_effective_hit_ratio", "target_effective_coverage",
    "binder_contact_neff_norm", "target_contact_neff_norm", "binder_contact_entropy", "target_contact_entropy",
    "binder_contact_span_norm", "target_contact_span", "target_contact_span_norm",
    "binder_field_active_roughness", "binder_field_active_smoothness", "binder_field_active_edges",
    "binder_field_core_roughness", "binder_field_core_smoothness", "binder_field_core_edges",
    "target_field_active_roughness", "target_field_active_smoothness", "target_field_active_edges",
    "target_field_core_roughness", "target_field_core_smoothness", "target_field_core_edges",
    "contact_map_edges", "contact_map_continuity_mean_dist", "contact_map_continuity_score", "contact_map_distortion_mean", "contact_map_jump_fraction",
    "target_linearity", "target_planarity", "target_compactness",
    "contacted_target_linearity", "contacted_target_planarity", "contacted_target_compactness",
    "microdomain_count", "micro_plane_alignment_weighted", "micro_center_distance_weighted", "micro_planarity_mean", "micro_fit_score_weighted",
    "desired_region_count", "desired_contact_count", "desired_effective_coverage", "desired_weight_fraction", "desired_mean_field",
    "desired_contact_neff_norm", "desired_contact_entropy", "desired_contact_span", "desired_contact_span_norm_global",
    "undesired_region_count", "undesired_contact_count", "undesired_effective_coverage", "undesired_weight_fraction", "undesired_mean_field",
    "undesired_contact_neff_norm", "undesired_contact_entropy", "undesired_contact_span", "undesired_contact_span_norm_global",
    "hotspot_region_count", "hotspot_contact_count", "hotspot_effective_coverage", "hotspot_weight_fraction", "hotspot_mean_field",
    "hotspot_contact_neff_norm", "hotspot_contact_entropy",
    "hotspot_neighborhood_radius", "hotspot_neighborhood_count", "hotspot_neighborhood_contact_count",
    "hotspot_neighborhood_effective_coverage", "hotspot_neighborhood_weight_fraction", "hotspot_neighborhood_mean_field",
    "hotspot_neighborhood_contact_neff_norm", "hotspot_neighborhood_contact_entropy",
    "hotspot_neighborhood_contact_span", "hotspot_neighborhood_contact_span_norm_global",
    "error",
]


def collect_pdbs(input_path: Optional[str], input_dir: Optional[str], max_files: Optional[int]) -> List[Path]:
    pdbs: List[Path] = []
    if input_path:
        pdbs.append(Path(input_path))
    if input_dir:
        pdbs.extend(sorted(Path(input_dir).glob("*.pdb")))
    if max_files is not None:
        pdbs = pdbs[:max_files]
    return pdbs


def write_metrics_csv(rows: List[Dict[str, Any]], output_csv: Path) -> None:
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    headers = list(FIELDNAMES)
    # Keep any unexpected columns as well.
    for r in rows:
        for k in r.keys():
            if k not in headers:
                headers.append(k)
    with open(output_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        for r in rows:
            writer.writerow(r)


def run_feature_extraction(args: argparse.Namespace) -> List[Dict[str, Any]]:
    desired_specs = parse_region_specs(args.desired_regions)
    undesired_specs = parse_region_specs(args.undesired_regions)
    hotspot_specs = parse_region_specs(args.hotspot_regions)
    pdbs = collect_pdbs(args.input, args.input_dir, args.max_files)
    if not pdbs:
        raise SystemExit("No PDB files found. Use --input or --input_dir, or provide --input_csv to skip extraction.")

    rows: List[Dict[str, Any]] = []
    for idx, pdb in enumerate(pdbs, start=1):
        try:
            row = compute_metrics(pdb, args, desired_specs, undesired_specs, hotspot_specs)
            row["error"] = ""
        except Exception as e:
            row = {"pdb_name": pdb.stem, "error": str(e)}
        rows.append(row)
        if idx % 500 == 0:
            print(f"Processed {idx}/{len(pdbs)} PDB files")
    return rows


# =============================================================================
# V4: scoring, filtering, ranking
# =============================================================================

def safe_float(x: Any, default: float = math.nan) -> float:
    if x is None:
        return default
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip()
    if s == "" or s.lower() in {"nan", "none", "null"}:
        return default
    try:
        return float(s)
    except Exception:
        return default


def is_valid_row(row: Dict[str, Any]) -> bool:
    e = str(row.get("error", "")).strip()
    return e == "" or e.lower() in {"nan", "none", "null"}


def finite_values(rows: List[Dict[str, Any]], col: str) -> List[float]:
    vals = []
    for r in rows:
        v = safe_float(r.get(col))
        if math.isfinite(v):
            vals.append(v)
    return vals


def quantile(vals: List[float], q: float) -> float:
    vals = sorted([v for v in vals if math.isfinite(v)])
    if not vals:
        return math.nan
    q = max(0.0, min(1.0, q))
    if len(vals) == 1:
        return vals[0]
    pos = q * (len(vals) - 1)
    lo = int(math.floor(pos))
    hi = int(math.ceil(pos))
    if lo == hi:
        return vals[lo]
    frac = pos - lo
    return vals[lo] * (1.0 - frac) + vals[hi] * frac


def median(vals: List[float]) -> float:
    return quantile(vals, 0.5)


def clip01(x: float) -> float:
    if not math.isfinite(x):
        return 0.5
    return max(0.0, min(1.0, x))


def robust_norm_value(x: float, lo: float, hi: float, higher_better: bool = True) -> float:
    if not math.isfinite(x) or not math.isfinite(lo) or not math.isfinite(hi) or abs(hi - lo) < EPS_RANK:
        return 0.5
    z = (x - lo) / (hi - lo)
    z = clip01(z)
    return z if higher_better else 1.0 - z


def normalize_weights(raw: Dict[str, float]) -> Dict[str, float]:
    s = sum(max(0.0, v) for v in raw.values())
    if s <= EPS_RANK:
        return {k: 0.0 for k in raw}
    return {k: max(0.0, v) / s for k, v in raw.items()}


def add_derived_metrics(rows: List[Dict[str, Any]]) -> None:
    for r in rows:
        raw8 = safe_float(r.get("raw_mean_shell_8"), 0.0)
        raw12 = safe_float(r.get("raw_mean_shell_12"), 0.0)
        if raw12 > EPS_RANK:
            r["shell_sensitivity_12_vs_8"] = (raw12 - raw8) / (raw12 + EPS_RANK)
            r["tight_contact_fraction_8_over_12"] = raw8 / (raw12 + EPS_RANK)
        else:
            r["shell_sensitivity_12_vs_8"] = 1.0
            r["tight_contact_fraction_8_over_12"] = 0.0

        eff = safe_float(r.get("effective_weight_sum"), 0.0)
        t_cov = safe_float(r.get("target_effective_coverage"), 0.0)
        t_len = safe_float(r.get("target_len"), math.nan)
        t_neff_norm = safe_float(r.get("target_contact_neff_norm"), math.nan)
        if math.isfinite(t_len) and math.isfinite(t_neff_norm):
            t_neff = max(t_neff_norm * t_len, EPS_RANK)
            r["contact_density_target_neff"] = eff / t_neff
        else:
            r["contact_density_target_neff"] = math.nan
        if math.isfinite(t_len):
            covered_count = max(t_cov * t_len, EPS_RANK)
            r["contact_density_target_covered"] = eff / covered_count
        else:
            r["contact_density_target_covered"] = math.nan

        center_dist = safe_float(r.get("micro_center_distance_weighted"), math.nan)
        if math.isfinite(center_dist):
            r["micro_center_distance_score"] = math.exp(-((center_dist - 7.5) ** 2) / (2.0 * 3.0 * 3.0))
        else:
            r["micro_center_distance_score"] = math.nan

        r["contacted_linearity_score"] = safe_float(r.get("contacted_target_linearity"), math.nan)
        r["contacted_planarity_score"] = safe_float(r.get("contacted_target_planarity"), math.nan)
        r["contacted_compactness_score"] = safe_float(r.get("contacted_target_compactness"), math.nan)

        # Helpful binary flags for reporting.
        r["has_desired_region"] = 1 if safe_float(r.get("desired_region_count"), 0.0) > 0 else 0
        r["has_undesired_region"] = 1 if safe_float(r.get("undesired_region_count"), 0.0) > 0 else 0
        r["has_hotspot_region"] = 1 if safe_float(r.get("hotspot_region_count"), 0.0) > 0 else 0
        r["has_hotspot_neighborhood"] = 1 if safe_float(r.get("hotspot_neighborhood_count"), 0.0) > 0 else 0


def infer_target_morphology(rows: List[Dict[str, Any]], alpha: float = 1.25, min_mode_weight: float = 0.05) -> Dict[str, float]:
    valid = [r for r in rows if is_valid_row(r)]
    L = median(finite_values(valid, "target_linearity"))
    P = median(finite_values(valid, "target_planarity"))
    C = median(finite_values(valid, "target_compactness"))
    if not math.isfinite(L): L = 0.0
    if not math.isfinite(P): P = 0.0
    if not math.isfinite(C): C = 0.0
    raw = {"line": max(L, 0.0) ** alpha, "plane": max(P, 0.0) ** alpha, "compact": max(C, 0.0) ** alpha}
    base = normalize_weights(raw)
    floor = max(0.0, min(min_mode_weight, 0.30))
    if 3.0 * floor >= 1.0:
        floor = 0.0
    mode = {k: floor + (1.0 - 3.0 * floor) * v for k, v in base.items()}
    mode = normalize_weights(mode)
    return {
        "target_linearity_median": L,
        "target_planarity_median": P,
        "target_compactness_median": C,
        "morph_alpha": alpha,
        "min_mode_weight": min_mode_weight,
        "mode_weight_line": mode["line"],
        "mode_weight_plane": mode["plane"],
        "mode_weight_compact": mode["compact"],
    }


def get_component_weights() -> Dict[str, Dict[str, Tuple[str, float, bool]]]:
    score_line = {
        "effective_contact_strength": ("effective_weight_sum", 0.17, True),
        "target_coverage": ("target_effective_coverage", 0.19, True),
        "target_span": ("target_contact_span_norm", 0.19, True),
        "target_neff": ("target_contact_neff_norm", 0.08, True),
        "contact_map_continuity": ("contact_map_continuity_score", 0.09, True),
        "binder_active_smooth": ("binder_field_active_smoothness", 0.08, True),
        "contacted_target_linearity": ("contacted_linearity_score", 0.06, True),
        "cb_closer": ("cb_closer_weight_ratio", 0.05, True),
        "front_facing": ("front_facing_weight_ratio", 0.04, True),
        "tight_contact_fraction": ("tight_contact_fraction_8_over_12", 0.04, True),
        "low_backfacing_cb_far": ("backfacing_cb_far_weight_ratio", 0.07, False),
        "low_contact_jump": ("contact_map_jump_fraction", 0.04, False),
        "low_shell_sensitivity": ("shell_sensitivity_12_vs_8", 0.03, False),
    }
    score_plane = {
        "micro_fit": ("micro_fit_score_weighted", 0.20, True),
        "micro_plane_alignment": ("micro_plane_alignment_weighted", 0.14, True),
        "micro_planarity": ("micro_planarity_mean", 0.10, True),
        "micro_center_distance": ("micro_center_distance_score", 0.08, True),
        "target_coverage": ("target_effective_coverage", 0.13, True),
        "effective_contact_strength": ("effective_weight_sum", 0.12, True),
        "contacted_target_planarity": ("contacted_planarity_score", 0.08, True),
        "contact_map_continuity": ("contact_map_continuity_score", 0.05, True),
        "cb_closer": ("cb_closer_weight_ratio", 0.04, True),
        "front_facing": ("front_facing_weight_ratio", 0.03, True),
        "low_backfacing_cb_far": ("backfacing_cb_far_weight_ratio", 0.07, False),
        "low_contact_jump": ("contact_map_jump_fraction", 0.03, False),
    }
    score_compact = {
        "effective_contact_strength": ("effective_weight_sum", 0.18, True),
        "contact_density_neff": ("contact_density_target_neff", 0.15, True),
        "target_coverage": ("target_effective_coverage", 0.10, True),
        "micro_fit": ("micro_fit_score_weighted", 0.10, True),
        "contacted_target_compactness": ("contacted_compactness_score", 0.07, True),
        "cb_closer": ("cb_closer_weight_ratio", 0.08, True),
        "front_facing": ("front_facing_weight_ratio", 0.06, True),
        "binder_core_smooth": ("binder_field_core_smoothness", 0.07, True),
        "binder_active_smooth": ("binder_field_active_smoothness", 0.04, True),
        "low_backfacing_cb_far": ("backfacing_cb_far_weight_ratio", 0.12, False),
        "low_contact_jump": ("contact_map_jump_fraction", 0.04, False),
        "low_shell_sensitivity": ("shell_sensitivity_12_vs_8", 0.04, False),
    }
    score_roughness = {
        "binder_active_smooth": ("binder_field_active_smoothness", 0.32, True),
        "binder_core_smooth": ("binder_field_core_smoothness", 0.24, True),
        "target_active_smooth": ("target_field_active_smoothness", 0.12, True),
        "target_core_smooth": ("target_field_core_smoothness", 0.08, True),
        "contact_map_continuity": ("contact_map_continuity_score", 0.14, True),
        "low_contact_jump": ("contact_map_jump_fraction", 0.05, False),
        "low_backfacing_cb_far": ("backfacing_cb_far_weight_ratio", 0.05, False),
    }
    score_microfit = {
        "micro_fit": ("micro_fit_score_weighted", 0.38, True),
        "micro_plane_alignment": ("micro_plane_alignment_weighted", 0.22, True),
        "micro_planarity": ("micro_planarity_mean", 0.14, True),
        "micro_center_distance": ("micro_center_distance_score", 0.10, True),
        "effective_contact_strength": ("effective_weight_sum", 0.05, True),
        "target_coverage": ("target_effective_coverage", 0.05, True),
        "low_backfacing_cb_far": ("backfacing_cb_far_weight_ratio", 0.04, False),
        "low_contact_jump": ("contact_map_jump_fraction", 0.02, False),
    }
    score_safety = {
        "low_backfacing_cb_far": ("backfacing_cb_far_weight_ratio", 0.24, False),
        "cb_closer": ("cb_closer_weight_ratio", 0.16, True),
        "low_contact_jump": ("contact_map_jump_fraction", 0.16, False),
        "binder_active_smooth": ("binder_field_active_smoothness", 0.14, True),
        "tight_contact_fraction": ("tight_contact_fraction_8_over_12", 0.10, True),
        "low_shell_sensitivity": ("shell_sensitivity_12_vs_8", 0.10, False),
        "front_facing": ("front_facing_weight_ratio", 0.06, True),
        "low_clash_pairs": ("clash_pairs", 0.04, False),
    }
    # New weak region score: useful when target is a large protein or hotspot/design region matters.
    score_region = {
        "desired_weight_fraction": ("desired_weight_fraction", 0.30, True),
        "desired_effective_coverage": ("desired_effective_coverage", 0.20, True),
        "desired_mean_field": ("desired_mean_field", 0.08, True),
        "desired_neff": ("desired_contact_neff_norm", 0.08, True),
        "hotspot_neighborhood_weight_fraction": ("hotspot_neighborhood_weight_fraction", 0.14, True),
        "hotspot_neighborhood_coverage": ("hotspot_neighborhood_effective_coverage", 0.08, True),
        "low_undesired_weight_fraction": ("undesired_weight_fraction", 0.08, False),
        "low_undesired_coverage": ("undesired_effective_coverage", 0.04, False),
    }
    # Separate hotspot-only diagnostic score. It is not directly a large weight in final score.
    score_hotspot = {
        "hotspot_neighborhood_weight_fraction": ("hotspot_neighborhood_weight_fraction", 0.40, True),
        "hotspot_neighborhood_coverage": ("hotspot_neighborhood_effective_coverage", 0.30, True),
        "hotspot_weight_fraction": ("hotspot_weight_fraction", 0.15, True),
        "hotspot_coverage": ("hotspot_effective_coverage", 0.15, True),
    }
    return {
        "score_line": score_line,
        "score_plane": score_plane,
        "score_compact": score_compact,
        "score_roughness": score_roughness,
        "score_microfit": score_microfit,
        "score_safety": score_safety,
        "score_region": score_region,
        "score_hotspot": score_hotspot,
    }


def collect_metric_specs(component_weights: Dict[str, Dict[str, Tuple[str, float, bool]]]) -> Dict[Tuple[str, bool], None]:
    specs = {}
    for comps in component_weights.values():
        for _, (col, _, higher) in comps.items():
            specs[(col, higher)] = None
    return specs


def compute_normalization_ranges(rows: List[Dict[str, Any]], specs: Dict[Tuple[str, bool], None],
                                 q_low: float, q_high: float) -> Dict[Tuple[str, bool], Tuple[float, float]]:
    valid = [r for r in rows if is_valid_row(r)]
    ranges = {}
    for col, higher in specs:
        vals = finite_values(valid, col)
        ranges[(col, higher)] = (quantile(vals, q_low), quantile(vals, q_high))
    return ranges


def add_weighted_score(row_index: int,
                       norm_cache: Dict[Tuple[int, str, bool], float],
                       components: Dict[str, Tuple[str, float, bool]]) -> float:
    total_w = sum(max(0.0, w) for _, w, _ in components.values())
    if total_w <= EPS_RANK:
        return 0.0
    score = 0.0
    for _, (col, w, higher) in components.items():
        score += (max(0.0, w) / total_w) * norm_cache.get((row_index, col, higher), 0.5)
    return clip01(score)


def region_is_available(rows: List[Dict[str, Any]]) -> bool:
    valid = [r for r in rows if is_valid_row(r)]
    for r in valid:
        if safe_float(r.get("desired_region_count"), 0.0) > 0:
            return True
        if safe_float(r.get("undesired_region_count"), 0.0) > 0:
            return True
        if safe_float(r.get("hotspot_region_count"), 0.0) > 0:
            return True
        if safe_float(r.get("hotspot_neighborhood_count"), 0.0) > 0:
            return True
    return False


def compute_scores(rows: List[Dict[str, Any]], morph: Dict[str, float],
                   q_low: float = 0.05, q_high: float = 0.95,
                   region_score_mode: str = "auto") -> Dict[str, Any]:
    comp_weights = get_component_weights()
    specs = collect_metric_specs(comp_weights)
    ranges = compute_normalization_ranges(rows, specs, q_low, q_high)

    norm_cache: Dict[Tuple[int, str, bool], float] = {}
    for i, r in enumerate(rows):
        for (col, higher), (lo, hi) in ranges.items():
            val = safe_float(r.get(col))
            norm_cache[(i, col, higher)] = robust_norm_value(val, lo, hi, higher_better=higher)

    region_available = region_is_available(rows)
    if region_score_mode == "off":
        use_region = False
    elif region_score_mode == "always":
        use_region = True
    else:
        use_region = region_available

    for i, r in enumerate(rows):
        if not is_valid_row(r):
            for score_name in list(comp_weights.keys()) + [
                "morphology_adaptive_score", "final_score_v4_original",
                "final_score_v4_region", "final_score_v4", "rank_consensus_score"
            ]:
                r[score_name] = math.nan
            continue

        for score_name, comps in comp_weights.items():
            r[score_name] = add_weighted_score(i, norm_cache, comps)

        adaptive = (
            morph["mode_weight_line"] * safe_float(r.get("score_line"), 0.0)
            + morph["mode_weight_plane"] * safe_float(r.get("score_plane"), 0.0)
            + morph["mode_weight_compact"] * safe_float(r.get("score_compact"), 0.0)
        )
        r["morphology_adaptive_score"] = clip01(adaptive)

        original = clip01(
            0.85 * safe_float(r.get("morphology_adaptive_score"), 0.0)
            + 0.10 * safe_float(r.get("score_safety"), 0.0)
            + 0.05 * safe_float(r.get("score_roughness"), 0.0)
        )
        r["final_score_v4_original"] = original

        if use_region:
            region_final = clip01(
                0.78 * safe_float(r.get("morphology_adaptive_score"), 0.0)
                + 0.10 * safe_float(r.get("score_safety"), 0.0)
                + 0.05 * safe_float(r.get("score_roughness"), 0.0)
                + 0.07 * safe_float(r.get("score_region"), 0.5)
            )
        else:
            region_final = original
        r["final_score_v4_region"] = region_final
        # Primary score for ranking; original score is preserved separately.
        r["final_score_v4"] = region_final

    add_rank_consensus(rows, morph)

    return {
        "component_weights": comp_weights,
        "normalization_ranges": ranges,
        "q_low": q_low,
        "q_high": q_high,
        "region_available": region_available,
        "region_score_mode": region_score_mode,
        "region_score_used": use_region,
    }


def add_rank_consensus(rows: List[Dict[str, Any]], morph: Dict[str, float]) -> None:
    valid = [r for r in rows if is_valid_row(r)]
    score_cols = [
        ("score_line", morph["mode_weight_line"]),
        ("score_plane", morph["mode_weight_plane"]),
        ("score_compact", morph["mode_weight_compact"]),
        ("score_roughness", 0.10),
        ("score_microfit", 0.05),
        ("score_safety", 0.10),
        ("score_region", 0.05),
    ]
    total = sum(w for _, w in score_cols)
    score_cols = [(c, w / total) for c, w in score_cols if total > 0]

    percentiles: Dict[str, Dict[str, float]] = {}
    for col, _ in score_cols:
        sorted_rows = sorted(valid, key=lambda r: safe_float(r.get(col), -1e99))
        n = len(sorted_rows)
        p = {}
        for idx, r in enumerate(sorted_rows):
            p[str(r.get("pdb_name", ""))] = idx / max(n - 1, 1)
        percentiles[col] = p

    for r in rows:
        if not is_valid_row(r):
            r["rank_consensus_score"] = math.nan
            continue
        name = str(r.get("pdb_name", ""))
        s = 0.0
        for col, w in score_cols:
            s += w * percentiles[col].get(name, 0.0)
        r["rank_consensus_score"] = clip01(s)


def filter_definition(region_filter: str = "off") -> Dict[str, Dict[str, Tuple[str, str, float]]]:
    definitions = {
        "broad": {
            "effective_weight_sum": ("effective_weight_sum", "min", 0.10),
            "target_effective_coverage": ("target_effective_coverage", "min", 0.10),
            "target_contact_span_norm": ("target_contact_span_norm", "min", 0.10),
            "backfacing_cb_far_weight_ratio": ("backfacing_cb_far_weight_ratio", "max", 0.95),
            "binder_field_active_roughness": ("binder_field_active_roughness", "max", 0.95),
            "contact_map_jump_fraction": ("contact_map_jump_fraction", "max", 0.95),
            "shell_sensitivity_12_vs_8": ("shell_sensitivity_12_vs_8", "max", 0.98),
            "clash_pairs": ("clash_pairs", "max", 0.99),
        },
        "medium": {
            "effective_weight_sum": ("effective_weight_sum", "min", 0.40),
            "target_effective_coverage": ("target_effective_coverage", "min", 0.40),
            "target_contact_span_norm": ("target_contact_span_norm", "min", 0.40),
            "backfacing_cb_far_weight_ratio": ("backfacing_cb_far_weight_ratio", "max", 0.85),
            "binder_field_active_roughness": ("binder_field_active_roughness", "max", 0.85),
            "contact_map_jump_fraction": ("contact_map_jump_fraction", "max", 0.85),
            "shell_sensitivity_12_vs_8": ("shell_sensitivity_12_vs_8", "max", 0.90),
            "cb_closer_weight_ratio": ("cb_closer_weight_ratio", "min", 0.20),
            "contact_map_continuity_score": ("contact_map_continuity_score", "min", 0.20),
            "clash_pairs": ("clash_pairs", "max", 0.99),
        },
        "strict": {
            "effective_weight_sum": ("effective_weight_sum", "min", 0.60),
            "target_effective_coverage": ("target_effective_coverage", "min", 0.60),
            "target_contact_span_norm": ("target_contact_span_norm", "min", 0.55),
            "backfacing_cb_far_weight_ratio": ("backfacing_cb_far_weight_ratio", "max", 0.75),
            "binder_field_active_roughness": ("binder_field_active_roughness", "max", 0.75),
            "contact_map_jump_fraction": ("contact_map_jump_fraction", "max", 0.75),
            "shell_sensitivity_12_vs_8": ("shell_sensitivity_12_vs_8", "max", 0.80),
            "cb_closer_weight_ratio": ("cb_closer_weight_ratio", "min", 0.35),
            "contact_map_continuity_score": ("contact_map_continuity_score", "min", 0.35),
            "score_safety": ("score_safety", "min", 0.35),
            "clash_pairs": ("clash_pairs", "max", 0.99),
        },
    }
    if region_filter in {"soft", "strict"}:
        # Soft region filters only remove clear outliers if region fields have variation.
        definitions["broad"].update({
            "score_region": ("score_region", "min", 0.05),
            "undesired_weight_fraction": ("undesired_weight_fraction", "max", 0.98),
        })
        definitions["medium"].update({
            "score_region": ("score_region", "min", 0.15),
            "undesired_weight_fraction": ("undesired_weight_fraction", "max", 0.95),
        })
        definitions["strict"].update({
            "score_region": ("score_region", "min", 0.25),
            "undesired_weight_fraction": ("undesired_weight_fraction", "max", 0.90),
        })
    if region_filter == "strict":
        definitions["strict"].update({
            "desired_weight_fraction": ("desired_weight_fraction", "min", 0.20),
        })
    return definitions


def compute_filter_thresholds(rows: List[Dict[str, Any]], region_filter: str = "off") -> Dict[str, Dict[str, Dict[str, Any]]]:
    valid = [r for r in rows if is_valid_row(r)]
    definitions = filter_definition(region_filter=region_filter)
    thresholds: Dict[str, Dict[str, Dict[str, Any]]] = {}
    for layer, rules in definitions.items():
        thresholds[layer] = {}
        for rule_name, (col, mode, q) in rules.items():
            vals = finite_values(valid, col)
            thr = quantile(vals, q)
            if col == "clash_pairs":
                thr = max(0.0, thr)
            thresholds[layer][rule_name] = {"column": col, "mode": mode, "quantile": q, "threshold": thr}
    return thresholds


def apply_layered_filters(rows: List[Dict[str, Any]], thresholds: Dict[str, Dict[str, Dict[str, Any]]]) -> None:
    for r in rows:
        if not is_valid_row(r):
            for layer in ["broad", "medium", "strict"]:
                r[f"filter_{layer}_pass"] = "NO"
                r[f"filter_{layer}_reasons"] = "error_not_empty"
            r["filter_level"] = "FAIL"
            continue
        for layer in ["broad", "medium", "strict"]:
            reasons = []
            for _, spec in thresholds[layer].items():
                col = spec["column"]
                mode = spec["mode"]
                thr = spec["threshold"]
                val = safe_float(r.get(col))
                if not math.isfinite(val) or not math.isfinite(thr):
                    continue
                if mode == "min" and val < thr:
                    reasons.append(f"low_{col}")
                elif mode == "max" and val > thr:
                    reasons.append(f"high_{col}")
            r[f"filter_{layer}_pass"] = "YES" if not reasons else "NO"
            r[f"filter_{layer}_reasons"] = ";".join(reasons)
        if r["filter_strict_pass"] == "YES":
            r["filter_level"] = "STRICT"
        elif r["filter_medium_pass"] == "YES":
            r["filter_level"] = "MEDIUM"
        elif r["filter_broad_pass"] == "YES":
            r["filter_level"] = "BROAD"
        else:
            r["filter_level"] = "FAIL"


def add_ranks(rows: List[Dict[str, Any]], score_cols: List[str]) -> None:
    valid = [r for r in rows if is_valid_row(r)]
    for col in score_cols:
        sorted_rows = sorted(valid, key=lambda r: safe_float(r.get(col), -1e99), reverse=True)
        for idx, r in enumerate(sorted_rows, start=1):
            r[f"rank_{col}"] = idx
        for r in rows:
            if not is_valid_row(r):
                r[f"rank_{col}"] = ""
    for layer in ["broad", "medium", "strict"]:
        passed = [r for r in valid if r.get(f"filter_{layer}_pass") == "YES"]
        sorted_passed = sorted(passed, key=lambda r: safe_float(r.get("final_score_v4"), -1e99), reverse=True)
        for idx, r in enumerate(sorted_passed, start=1):
            r[f"rank_final_in_{layer}"] = idx
        for r in rows:
            if r.get(f"filter_{layer}_pass") != "YES":
                r[f"rank_final_in_{layer}"] = ""


def names_match(row_name: str, ref: str) -> bool:
    row_name = str(row_name).strip()
    ref = str(ref).strip().replace(".pdb", "")
    if not row_name or not ref:
        return False
    return row_name == ref or row_name.endswith(ref)


def find_references(rows: List[Dict[str, Any]], refs: List[str]) -> List[Dict[str, Any]]:
    out = []
    for ref in refs:
        found = None
        for r in rows:
            if names_match(str(r.get("pdb_name", "")), ref):
                found = r
                break
        if found is None:
            out.append({"_reference_query": ref, "pdb_name": "NOT_FOUND", "filter_level": "", "filter_strict_reasons": "reference_not_found"})
        else:
            item = dict(found)
            item["_reference_query"] = ref
            out.append(item)
    return out


# =============================================================================
# Output helpers
# =============================================================================

DISPLAY_COLS = [
    "pdb_name", "filter_level", "filter_broad_pass", "filter_broad_reasons",
    "filter_medium_pass", "filter_medium_reasons", "filter_strict_pass", "filter_strict_reasons",
    "final_score_v4", "final_score_v4_original", "final_score_v4_region",
    "morphology_adaptive_score", "rank_consensus_score",
    "score_line", "score_plane", "score_compact", "score_roughness", "score_microfit", "score_safety",
    "score_region", "score_hotspot",
    "rank_final_score_v4", "rank_final_score_v4_original", "rank_final_score_v4_region",
    "rank_final_in_broad", "rank_final_in_medium", "rank_final_in_strict",
    "effective_weight_sum", "target_effective_coverage", "target_contact_span_norm", "target_contact_neff_norm", "target_contact_entropy",
    "desired_region_count", "desired_contact_count", "desired_effective_coverage", "desired_weight_fraction", "desired_mean_field",
    "desired_contact_neff_norm", "desired_contact_span_norm_global",
    "hotspot_region_count", "hotspot_contact_count", "hotspot_effective_coverage", "hotspot_weight_fraction", "hotspot_mean_field",
    "hotspot_neighborhood_count", "hotspot_neighborhood_contact_count", "hotspot_neighborhood_effective_coverage",
    "hotspot_neighborhood_weight_fraction", "hotspot_neighborhood_mean_field",
    "undesired_region_count", "undesired_contact_count", "undesired_effective_coverage", "undesired_weight_fraction", "undesired_mean_field",
    "raw_mean_shell_8", "raw_mean_shell_10", "raw_mean_shell_12", "tight_contact_fraction_8_over_12", "shell_sensitivity_12_vs_8",
    "front_facing_weight_ratio", "tangential_weight_ratio", "backfacing_cb_far_weight_ratio", "cb_closer_weight_ratio", "cb_delta_weighted_mean",
    "binder_field_active_roughness", "binder_field_active_smoothness", "binder_field_core_roughness", "binder_field_core_smoothness",
    "target_field_active_roughness", "target_field_active_smoothness", "target_field_core_roughness", "target_field_core_smoothness",
    "contact_map_continuity_score", "contact_map_jump_fraction",
    "micro_plane_alignment_weighted", "micro_center_distance_weighted", "micro_center_distance_score", "micro_planarity_mean", "micro_fit_score_weighted",
    "target_linearity", "target_planarity", "target_compactness",
    "contacted_target_linearity", "contacted_target_planarity", "contacted_target_compactness",
    "clash_pairs", "min_distance_global", "error",
]


def fmt(v: Any) -> Any:
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return ""
    return v


def autosize_sheet(ws, max_width: int = 42) -> None:
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col[:200]:
            try:
                max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), max_width)


def write_sheet(ws, rows: List[Dict[str, Any]], headers: List[str]) -> None:
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append([fmt(r.get(h, "")) for h in headers])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    autosize_sheet(ws)


def write_kv_sheet(ws, title: str, rows: List[Tuple[Any, Any, Any]]) -> None:
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["key", "value", "note"])
    for cell in ws[3]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = "A4"
    autosize_sheet(ws)


def metric_notes() -> List[Dict[str, str]]:
    rows = [
        ("final_score_v4", "Primary score", "By default equals region-aware final score if region/hotspot exists; otherwise equals original score"),
        ("final_score_v4_original", "Original V4 score", "0.85 adaptive + 0.10 safety + 0.05 roughness; preserved for comparison"),
        ("final_score_v4_region", "Region-aware score", "0.78 adaptive + 0.10 safety + 0.05 roughness + 0.07 region if region is available"),
        ("score_region", "Weak design-region score", "Desired/hotspot positive, undesired negative; intentionally weak"),
        ("score_hotspot", "Hotspot diagnostic score", "Hotspot and hotspot-neighborhood contact quality; not a dominant final term"),
        ("morphology_adaptive_score", "Dynamic morphology score", "line/plane/compact mixed by target morphology"),
        ("score_line", "Linear-target mode", "For elongated targets; emphasizes coverage, span, continuity, active smoothness"),
        ("score_plane", "Plane-target mode", "For flat surfaces; emphasizes microdomain plane complementarity"),
        ("score_compact", "Compact-target mode", "For pockets/local patches; emphasizes dense and safe local contact"),
        ("score_safety", "Safety stabilizer", "Backfacing, jumps, shell sensitivity, clash-related stabilizer"),
        ("desired_weight_fraction", "Desired-region hit", "Fraction of total target contact weight falling in desired region"),
        ("desired_effective_coverage", "Desired-region coverage", "Fraction of desired-region residues with contact field > threshold"),
        ("hotspot_neighborhood_weight_fraction", "Hotspot-neighborhood hit", "Fraction of total target contact weight in hotspot neighborhood"),
        ("undesired_weight_fraction", "Undesired-region leakage", "Fraction of total target contact weight falling in undesired region; lower is better"),
        ("target_effective_coverage", "Global target coverage", "Fraction of all target residues effectively contacted; if target is full large protein, interpret carefully"),
        ("target_contact_span_norm", "Global target span", "Spatial spread of contacted target region normalized by full target span"),
        ("effective_weight_sum", "Effective contact amount", "Sum of soft contact matrix W"),
        ("contact_map_continuity_score", "Contact-map continuity", "Adjacent binder residues map to nearby target regions"),
        ("shell_sensitivity_12_vs_8", "Shell sensitivity", "High means loose 12 Å shell contributes much more than tight 8 Å shell"),
    ]
    return [{"metric": a, "meaning": b, "note": c} for a, b, c in rows]


def build_weight_filter_rows(morph: Dict[str, float], thresholds: Dict[str, Dict[str, Dict[str, Any]]],
                             metadata: Dict[str, Any], args: argparse.Namespace) -> List[Tuple[Any, Any, Any]]:
    out: List[Tuple[Any, Any, Any]] = []
    out.append(("RUN_CONFIG", "", "manual parameters used by this run"))
    for key in ["input", "input_dir", "input_csv", "binder_chain", "desired_regions", "undesired_regions", "hotspot_regions", "hotspot_expand_radius", "region_score_mode", "region_filter"]:
        out.append((key, getattr(args, key, ""), ""))
    out.append(("region_available", metadata.get("region_available"), "derived from region/hotspot count columns"))
    out.append(("region_score_used", metadata.get("region_score_used"), "whether final_score_v4 includes score_region"))
    out.append(("", "", ""))
    out.append(("TARGET_MORPHOLOGY", "", ""))
    out.append(("target_linearity_median", morph["target_linearity_median"], "median across valid rows"))
    out.append(("target_planarity_median", morph["target_planarity_median"], "median across valid rows"))
    out.append(("target_compactness_median", morph["target_compactness_median"], "median across valid rows"))
    out.append(("mode_weight_line", morph["mode_weight_line"], "line-target mode weight"))
    out.append(("mode_weight_plane", morph["mode_weight_plane"], "plane-target mode weight"))
    out.append(("mode_weight_compact", morph["mode_weight_compact"], "compact-target mode weight"))
    out.append(("", "", ""))
    out.append(("FINAL_SCORE_ORIGINAL", "0.85*morphology_adaptive + 0.10*safety + 0.05*roughness", "preserved as final_score_v4_original"))
    out.append(("FINAL_SCORE_REGION", "0.78*morphology_adaptive + 0.10*safety + 0.05*roughness + 0.07*score_region", "used if region_score_used=True"))
    out.append(("", "", ""))
    out.append(("FILTER_THRESHOLDS", "", "percentile-based thresholds"))
    for layer in ["broad", "medium", "strict"]:
        out.append((f"{layer.upper()}_FILTER", "", ""))
        for rule_name, spec in thresholds[layer].items():
            out.append((f"{layer}.{rule_name}", spec["threshold"], f"{spec['mode']} {spec['column']} at q={spec['quantile']}"))
    out.append(("", "", ""))
    out.append(("COMPONENT_WEIGHTS", "", "weights inside each score; normalized within score"))
    comp_weights = metadata["component_weights"]
    for score_name, comps in comp_weights.items():
        total = sum(w for _, w, _ in comps.values())
        out.append((score_name, "", ""))
        for comp_name, (col, w, higher) in comps.items():
            direction = "higher_better" if higher else "lower_better"
            out.append((f"{score_name}.{comp_name}", w / total if total > EPS_RANK else w, f"{col}; {direction}"))
    out.append(("", "", ""))
    out.append(("NORMALIZATION", "", f"robust percentile q={metadata['q_low']} to q={metadata['q_high']}"))
    for (col, higher), (lo, hi) in metadata["normalization_ranges"].items():
        direction = "higher_better" if higher else "lower_better"
        out.append((f"range.{col}.{direction}", f"{lo} to {hi}", ""))
    return out


def manual_config_notes(args: argparse.Namespace) -> List[Dict[str, str]]:
    rows = [
        {"item": "原始骨架路径", "how_to_modify": "--input_dir ./pdbs 或 --input ./one.pdb；已有CSV时用 --input_csv", "current": str(args.input_dir or args.input_csv or args.input or "")},
        {"item": "binder链", "how_to_modify": "--binder_chain A；如果binder不是A链就改成对应链ID", "current": str(args.binder_chain)},
        {"item": "design/desired region", "how_to_modify": "--desired_regions \"B:45-60,B:80-95\"；完整大蛋白建议一定填写", "current": str(args.desired_regions)},
        {"item": "undesired region", "how_to_modify": "--undesired_regions \"B:150-180\"；没有就留空", "current": str(args.undesired_regions)},
        {"item": "hotspot", "how_to_modify": "--hotspot_regions \"B:48,B:52\"；建议填RFdiffusion生成时使用的hotspot", "current": str(args.hotspot_regions)},
        {"item": "hotspot邻域半径", "how_to_modify": "--hotspot_expand_radius 8.0；想更宽松可10.0，更严格可6.0", "current": str(args.hotspot_expand_radius)},
        {"item": "region分数模式", "how_to_modify": "--region_score_mode auto/off/always；auto表示有region/hotspot才加入弱加分", "current": str(args.region_score_mode)},
        {"item": "region过滤", "how_to_modify": "--region_filter off/soft/strict；默认off，避免破坏旧筛选池", "current": str(args.region_filter)},
    ]
    return rows


def write_excel(output_xlsx: Path, raw_original: List[Dict[str, Any]], rows: List[Dict[str, Any]], metadata: Dict[str, Any],
                morph: Dict[str, float], thresholds: Dict[str, Dict[str, Dict[str, Any]]], references: List[Dict[str, Any]],
                args: argparse.Namespace, top_n_sheets: int = 500) -> None:
    if Workbook is None:
        print("WARNING: openpyxl not installed; skipping Excel output.", file=sys.stderr)
        return
    wb = Workbook()
    wb.remove(wb.active)
    raw_headers = list(raw_original[0].keys()) if raw_original else []
    headers: List[str] = []
    for h in DISPLAY_COLS:
        if (rows and h in rows[0]) or h in raw_headers:
            headers.append(h)
    if rows:
        for h in list(rows[0].keys()):
            if h not in headers:
                headers.append(h)

    ws = wb.create_sheet("01_v3_metrics_raw")
    write_sheet(ws, raw_original, raw_headers if raw_headers else headers)
    ws = wb.create_sheet("02_scored_all")
    write_sheet(ws, rows, headers)

    valid = [r for r in rows if is_valid_row(r)]
    for layer, sheet in [("broad", "03_filter_broad"), ("medium", "04_filter_medium"), ("strict", "05_filter_strict")]:
        ws = wb.create_sheet(sheet)
        write_sheet(ws, [r for r in valid if r.get(f"filter_{layer}_pass") == "YES"], headers)

    def ranked_by(col: str, pool: Optional[List[Dict[str, Any]]] = None) -> List[Dict[str, Any]]:
        p = pool if pool is not None else valid
        return sorted(p, key=lambda r: safe_float(r.get(col), -1e99), reverse=True)[:top_n_sheets]

    sheet_defs = [
        ("06_ranked_primary", "final_score_v4", valid),
        ("07_ranked_strict", "final_score_v4", [r for r in valid if r.get("filter_strict_pass") == "YES"]),
        ("08_ranked_original", "final_score_v4_original", valid),
        ("09_ranked_region", "final_score_v4_region", valid),
        ("10_rank_line", "score_line", valid),
        ("11_rank_plane", "score_plane", valid),
        ("12_rank_compact", "score_compact", valid),
        ("13_rank_roughness", "score_roughness", valid),
        ("14_rank_microfit", "score_microfit", valid),
        ("15_rank_region", "score_region", valid),
        ("16_rank_hotspot", "score_hotspot", valid),
    ]
    for sheet, col, pool in sheet_defs:
        ws = wb.create_sheet(sheet)
        write_sheet(ws, ranked_by(col, pool), headers)

    ws = wb.create_sheet("17_references")
    ref_headers = ["_reference_query"] + headers
    write_sheet(ws, references, ref_headers)

    ws = wb.create_sheet("18_weights_filters")
    write_kv_sheet(ws, "Integrated contact-field ranking weights, filters, and run config",
                   build_weight_filter_rows(morph, thresholds, metadata, args))

    ws = wb.create_sheet("19_metric_notes")
    write_sheet(ws, metric_notes(), ["metric", "meaning", "note"])

    ws = wb.create_sheet("20_manual_config_notes")
    write_sheet(ws, manual_config_notes(args), ["item", "how_to_modify", "current"])

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx)


def write_scored_csv(rows: List[Dict[str, Any]], output_csv: Path) -> None:
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    headers: List[str] = []
    for h in DISPLAY_COLS:
        if rows and h in rows[0]:
            headers.append(h)
    if rows:
        for k in rows[0].keys():
            if k not in headers:
                headers.append(k)
    with open(output_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        for r in rows:
            writer.writerow(r)


def write_txt_report(output_txt: Path, rows: List[Dict[str, Any]], morph: Dict[str, float],
                     thresholds: Dict[str, Dict[str, Dict[str, Any]]], references: List[Dict[str, Any]],
                     metadata: Dict[str, Any], args: argparse.Namespace, top_k: int = 30) -> None:
    valid = [r for r in rows if is_valid_row(r)]
    pools = {
        "all": valid,
        "broad": [r for r in valid if r.get("filter_broad_pass") == "YES"],
        "medium": [r for r in valid if r.get("filter_medium_pass") == "YES"],
        "strict": [r for r in valid if r.get("filter_strict_pass") == "YES"],
    }
    lines: List[str] = []
    lines.append("Integrated V3 + V4 region/hotspot-aware ranking report")
    lines.append("=" * 78)
    lines.append("")
    lines.append("Run config:")
    lines.append(f"  input/input_dir/input_csv = {args.input or args.input_dir or args.input_csv}")
    lines.append(f"  binder_chain              = {args.binder_chain}")
    lines.append(f"  desired_regions           = {args.desired_regions}")
    lines.append(f"  undesired_regions         = {args.undesired_regions}")
    lines.append(f"  hotspot_regions           = {args.hotspot_regions}")
    lines.append(f"  hotspot_expand_radius     = {args.hotspot_expand_radius}")
    lines.append(f"  region_score_mode         = {args.region_score_mode}")
    lines.append(f"  region_score_used         = {metadata.get('region_score_used')}")
    lines.append(f"  region_filter             = {args.region_filter}")
    lines.append("")
    lines.append("Target morphology:")
    lines.append(f"  target_linearity_median   = {morph['target_linearity_median']:.6f}")
    lines.append(f"  target_planarity_median   = {morph['target_planarity_median']:.6f}")
    lines.append(f"  target_compactness_median = {morph['target_compactness_median']:.6f}")
    lines.append("")
    lines.append("Dynamic mode weights:")
    lines.append(f"  line    = {morph['mode_weight_line']:.6f}")
    lines.append(f"  plane   = {morph['mode_weight_plane']:.6f}")
    lines.append(f"  compact = {morph['mode_weight_compact']:.6f}")
    lines.append("")
    lines.append("Primary scores:")
    lines.append("  final_score_v4_original = 0.85*morphology_adaptive_score + 0.10*score_safety + 0.05*score_roughness")
    lines.append("  final_score_v4_region   = 0.78*morphology_adaptive_score + 0.10*score_safety + 0.05*score_roughness + 0.07*score_region")
    lines.append("  final_score_v4          = primary ranking score; equals region-aware score if region_score_used=True")
    lines.append("")
    lines.append("Pool sizes:")
    for k, v in pools.items():
        lines.append(f"  {k:>6}: {len(v)}")
    lines.append("")
    lines.append("Filter thresholds:")
    for layer in ["broad", "medium", "strict"]:
        lines.append(f"  [{layer}]")
        for rule_name, spec in thresholds[layer].items():
            lines.append(f"    {rule_name}: {spec['mode']} {spec['column']} {spec['threshold']}  (q={spec['quantile']})")
    lines.append("")

    cols = [
        "pdb_name", "filter_level", "final_score_v4", "final_score_v4_original", "final_score_v4_region",
        "morphology_adaptive_score", "score_line", "score_safety", "score_roughness", "score_region", "score_hotspot",
        "effective_weight_sum", "target_effective_coverage", "target_contact_span_norm",
        "desired_weight_fraction", "desired_effective_coverage",
        "hotspot_neighborhood_weight_fraction", "hotspot_neighborhood_effective_coverage",
        "undesired_weight_fraction",
    ]
    for pool_name in ["all", "broad", "medium", "strict"]:
        ranked = sorted(pools[pool_name], key=lambda r: safe_float(r.get("final_score_v4"), -1e99), reverse=True)
        lines.append(f"Top {min(top_k, len(ranked))} in {pool_name} pool by final_score_v4:")
        for idx, r in enumerate(ranked[:top_k], start=1):
            parts = []
            for c in cols:
                v = r.get(c, "")
                if isinstance(v, float):
                    parts.append(f"{c}={v:.5f}")
                else:
                    parts.append(f"{c}={v}")
            lines.append(f"  #{idx}: " + " | ".join(parts))
        lines.append("")
    lines.append("Reference candidates:")
    if references:
        for r in references:
            q = r.get("_reference_query", "")
            name = r.get("pdb_name", "NOT_FOUND")
            if name == "NOT_FOUND":
                lines.append(f"  {q}: NOT_FOUND")
            else:
                lines.append(
                    f"  {q} -> {name}: "
                    f"rank_primary={r.get('rank_final_score_v4','')}, "
                    f"rank_original={r.get('rank_final_score_v4_original','')}, "
                    f"rank_region={r.get('rank_final_score_v4_region','')}, "
                    f"rank_strict={r.get('rank_final_in_strict','')}, "
                    f"level={r.get('filter_level','')}, "
                    f"final={safe_float(r.get('final_score_v4')):.5f}, "
                    f"region={safe_float(r.get('score_region')):.5f}, "
                    f"hotspot={safe_float(r.get('score_hotspot')):.5f}, "
                    f"strict_reasons={r.get('filter_strict_reasons','')}"
                )
    else:
        lines.append("  None")
    lines.append("")
    lines.append("Interpretation reminders:")
    lines.append("  - final_score_v4_original preserves the previous V4 logic.")
    lines.append("  - score_region is intentionally weak; it should prevent target/design-region drift, not dominate geometry quality.")
    lines.append("  - If target is a full large protein, desired/hotspot-region metrics are more interpretable than global target coverage.")
    lines.append("  - Broad / medium / strict are candidate pools, not ground truth.")
    lines.append("  - MPNN + AF2/Rosetta/experimental validation is still required downstream.")
    lines.append("")
    output_txt.parent.mkdir(parents=True, exist_ok=True)
    output_txt.write_text("\n".join(lines), encoding="utf-8")


def read_csv_rows(path: Path) -> List[Dict[str, Any]]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        return [dict(r) for r in reader]


def run_scoring_pipeline(rows: List[Dict[str, Any]], args: argparse.Namespace) -> Tuple[List[Dict[str, Any]], Dict[str, Any], Dict[str, float], Dict[str, Dict[str, Dict[str, Any]]], List[Dict[str, Any]]]:
    raw_original = [dict(r) for r in rows]
    add_derived_metrics(rows)
    morph = infer_target_morphology(rows, alpha=args.morph_alpha, min_mode_weight=args.min_mode_weight)
    metadata = compute_scores(rows, morph, q_low=args.q_low, q_high=args.q_high, region_score_mode=args.region_score_mode)
    thresholds = compute_filter_thresholds(rows, region_filter=args.region_filter)
    apply_layered_filters(rows, thresholds)
    rank_cols = [
        "final_score_v4", "final_score_v4_original", "final_score_v4_region",
        "morphology_adaptive_score", "rank_consensus_score",
        "score_line", "score_plane", "score_compact", "score_roughness", "score_microfit", "score_safety", "score_region", "score_hotspot",
    ]
    add_ranks(rows, rank_cols)
    refs = [x.strip() for x in str(args.reference_names).split(",") if x.strip()]
    ref_rows = find_references(rows, refs)
    return raw_original, metadata, morph, thresholds, ref_rows


def build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Integrated V3 contact metrics + V4 region/hotspot-aware ranking.")
    # Input/output.
    p.add_argument("--input", type=str, default=None, help="Single PDB file. Optional if --input_dir or --input_csv is used.")
    p.add_argument("--input_dir", type=str, default=None, help="Directory containing PDB files.")
    p.add_argument("--input_csv", type=str, default=None, help="Existing V3 metrics CSV; skips PDB feature extraction.")
    p.add_argument("--output_prefix", type=str, required=True, help="Output prefix, e.g. ./results/qbqc_integrated")
    p.add_argument("--max_files", type=int, default=None, help="Debug: process only first N PDB files.")

    # Manual region settings.
    p.add_argument("--binder_chain", type=str, default="A", help="Binder chain ID. Target = all non-binder chains.")
    p.add_argument("--desired_regions", type=str, default="", help="Design/desired target regions, e.g. B:45-60,B:80.")
    p.add_argument("--undesired_regions", type=str, default="", help="Undesired target regions, e.g. B:150-180.")
    p.add_argument("--hotspot_regions", type=str, default="", help="Hotspot residues/regions, e.g. B:48,B:52,B:56.")
    p.add_argument("--hotspot_expand_radius", type=float, default=8.0, help="CA-distance radius to form hotspot neighborhood.")

    # V3 geometry/contact parameters.
    p.add_argument("--shell_max", type=float, default=10.0)
    p.add_argument("--ideal_dist", type=float, default=7.5)
    p.add_argument("--density_sigma", type=float, default=2.0)
    p.add_argument("--clash_cutoff", type=float, default=3.5)
    p.add_argument("--graph_cutoff", type=float, default=10.0)
    p.add_argument("--contact_threshold", type=float, default=0.1)
    p.add_argument("--tangent_floor", type=float, default=0.45)
    p.add_argument("--front_dot", type=float, default=0.25)
    p.add_argument("--back_dot", type=float, default=-0.25)
    p.add_argument("--dot_sigma", type=float, default=0.20)
    p.add_argument("--back_penalty", type=float, default=0.25)
    p.add_argument("--cb_far_delta", type=float, default=0.8)
    p.add_argument("--cb_delta_sigma", type=float, default=0.5)
    p.add_argument("--cb_near_cutoff", type=float, default=10.0)
    p.add_argument("--cb_support_cutoff", type=float, default=12.0)
    p.add_argument("--cb_sigma", type=float, default=1.5)
    p.add_argument("--micro_radius", type=float, default=8.0)
    p.add_argument("--micro_dist_sigma", type=float, default=3.0)
    p.add_argument("--max_microdomains", type=int, default=80)
    p.add_argument("--jump_cutoff", type=float, default=15.0)

    # V4 ranking parameters.
    p.add_argument("--reference_names", default="", help="Comma-separated reference candidates for sanity check.")
    p.add_argument("--top_n_sheets", type=int, default=500, help="Max rows in each ranked Excel sheet.")
    p.add_argument("--top_k_report", type=int, default=30, help="Top K shown in TXT report.")
    p.add_argument("--morph_alpha", type=float, default=1.25, help="Power for morphology-to-mode conversion.")
    p.add_argument("--min_mode_weight", type=float, default=0.05, help="Minimum retained weight for each morphology mode.")
    p.add_argument("--q_low", type=float, default=0.05, help="Low quantile for robust normalization.")
    p.add_argument("--q_high", type=float, default=0.95, help="High quantile for robust normalization.")
    p.add_argument("--region_score_mode", choices=["auto", "off", "always"], default="auto",
                   help="auto: use weak region score only if region/hotspot exists; off: original final score only; always: always include region score.")
    p.add_argument("--region_filter", choices=["off", "soft", "strict"], default="off",
                   help="off preserves old filters; soft/strict add region-related percentile filters.")
    return p


def main() -> None:
    args = build_arg_parser().parse_args()
    prefix = Path(args.output_prefix)
    metrics_csv = prefix.with_name(prefix.name + "_metrics.csv")
    scored_csv = prefix.with_name(prefix.name + "_scored.csv")
    output_xlsx = prefix.with_name(prefix.name + "_ranking.xlsx")
    output_txt = prefix.with_name(prefix.name + "_report.txt")

    if args.input_csv:
        rows = read_csv_rows(Path(args.input_csv))
        if not rows:
            raise SystemExit("No rows found in input CSV.")
        # Keep a copy as metrics CSV with the new output name for reproducibility.
        write_metrics_csv(rows, metrics_csv)
        print(f"Loaded existing metrics CSV: {args.input_csv}")
    else:
        rows = run_feature_extraction(args)
        write_metrics_csv(rows, metrics_csv)
        print(f"Wrote metrics CSV: {metrics_csv}")

    raw_original, metadata, morph, thresholds, ref_rows = run_scoring_pipeline(rows, args)
    write_scored_csv(rows, scored_csv)
    write_excel(output_xlsx, raw_original, rows, metadata, morph, thresholds, ref_rows, args, top_n_sheets=args.top_n_sheets)
    write_txt_report(output_txt, rows, morph, thresholds, ref_rows, metadata, args, top_k=args.top_k_report)

    valid = [r for r in rows if is_valid_row(r)]
    broad = [r for r in valid if r.get("filter_broad_pass") == "YES"]
    medium = [r for r in valid if r.get("filter_medium_pass") == "YES"]
    strict = [r for r in valid if r.get("filter_strict_pass") == "YES"]
    print(f"Wrote scored CSV: {scored_csv}")
    print(f"Wrote Excel workbook: {output_xlsx}")
    print(f"Wrote TXT report: {output_txt}")
    print("")
    print("Target morphology:")
    print(f"  linearity   = {morph['target_linearity_median']:.6f}")
    print(f"  planarity   = {morph['target_planarity_median']:.6f}")
    print(f"  compactness = {morph['target_compactness_median']:.6f}")
    print("Mode weights:")
    print(f"  line    = {morph['mode_weight_line']:.6f}")
    print(f"  plane   = {morph['mode_weight_plane']:.6f}")
    print(f"  compact = {morph['mode_weight_compact']:.6f}")
    print("Region scoring:")
    print(f"  region_available = {metadata.get('region_available')}")
    print(f"  region_score_used = {metadata.get('region_score_used')}")
    print("")
    print("Pool sizes:")
    print(f"  all    = {len(valid)}")
    print(f"  broad  = {len(broad)}")
    print(f"  medium = {len(medium)}")
    print(f"  strict = {len(strict)}")
    ranked = sorted(strict if strict else medium if medium else broad if broad else valid,
                    key=lambda r: safe_float(r.get("final_score_v4"), -1e99), reverse=True)
    print("")
    print("Top candidates in best available pool:")
    for idx, r in enumerate(ranked[:10], start=1):
        print(
            f"  #{idx:02d} {r.get('pdb_name')} "
            f"level={r.get('filter_level')} "
            f"final={safe_float(r.get('final_score_v4')):.5f} "
            f"orig={safe_float(r.get('final_score_v4_original')):.5f} "
            f"region={safe_float(r.get('score_region')):.5f} "
            f"hotspot={safe_float(r.get('score_hotspot')):.5f} "
            f"coverage={safe_float(r.get('target_effective_coverage')):.5f} "
            f"desired_frac={safe_float(r.get('desired_weight_fraction')):.5f} "
            f"undesired_frac={safe_float(r.get('undesired_weight_fraction')):.5f}"
        )


if __name__ == "__main__":
    main()
